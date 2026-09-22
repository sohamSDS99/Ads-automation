"""The `brand_book` connector (Stage 03 PRD §10.2).

A brand book is the one input to this stage that arrives as a binary, and law 30
says a binary never enters a prompt. So this connector's output is not "the
document" — it is **spans with provenance**: a piece of text, the page it sat
on, and the box it occupied. The bytes stay on the Volume behind `asset_path`
and nothing downstream ever holds them.

Provenance is the product here, not a nicety. G5 asks a brand owner to confirm
rules that were derived from their own document; "clear space is 1× the logo
height (page 4, top right)" is checkable and "clear space is 1× the logo height"
is a thing they have to take on faith.

Each format gives up its geometry differently, and pretending otherwise would
mean inventing coordinates:

* **PDF** has no per-run box. `pypdf` reports the text matrix through a visitor
  callback, so the origin is real and the extent is estimated from the font
  size and the character count. The estimate is marked as such — see `bbox`.
* **PPTX** carries a true box per shape in its own object model, in EMU.
* **DOCX** has no geometry at all until Word lays it out, so a paragraph index
  is the honest answer and `bbox` is `None`.
"""

from __future__ import annotations

import io
from typing import Any

import structlog

from agent.connectors.base import ConnectorDegraded, EvidenceDraft, ReadOnlyConnector
from agent.db.models import EvidenceSource

log = structlog.get_logger(__name__)

#: The evidence `kind` every span carries (PRD §10.2). One spelling, imported by
#: the connector that writes it and by the three nodes that read it.
SPAN = "brand_book_span"

#: An embedded or loose image, kept as a *candidate* logo. Which of them is the
#: logo is 3.1.3's judgement and G5's confirmation, not this connector's.
ASSET = "brand_book_asset"

#: One colour token with the page it was sampled from.
COLOUR = "brand_book_colour"

IMAGE_SUFFIXES = frozenset({"png", "jpg", "jpeg", "gif", "webp", "bmp", "tif", "tiff"})

#: Helvetica's average advance width as a fraction of the point size. Used only
#: to give a span a plausible *extent*; the origin is exact.
AVERAGE_GLYPH_WIDTH = 0.5


class BrandBookConnector(ReadOnlyConnector):
    """Brand books in, spans out. Never the other way round.

    Read-only by nature — it parses an upload — and marked as such so that a
    node naming it passes the executor's `assert_read_only` check.
    """

    name = "brand_book"
    #: PRD §10.2 names `source='web'` explicitly. `UPLOAD` would arguably be the
    #: better label — migration 0008 added it precisely so a pricing PDF does
    #: not appear under a heading that means something else — but the PRD is
    #: specific and Stage 04 reads these rows by source. Flagged as an open
    #: question rather than changed unilaterally.
    source = EvidenceSource.WEB

    def extract(self, filename: str, content: bytes, *, asset_path: str) -> list[EvidenceDraft]:
        """One file to spans. Synchronous: nothing here touches the network."""
        suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if suffix == "pdf":
            return self._pdf(content, asset_path=asset_path)
        if suffix == "docx":
            return self._docx(content, asset_path=asset_path)
        if suffix == "pptx":
            return self._pptx(content, asset_path=asset_path)
        if suffix in IMAGE_SUFFIXES:
            return self._image(content, asset_path=asset_path)
        raise ConnectorDegraded(f"{filename}: not a brand-book format this connector reads")

    def _pdf(self, content: bytes, *, asset_path: str) -> list[EvidenceDraft]:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        drafts: list[EvidenceDraft] = []
        pages = 0
        for number, page in enumerate(reader.pages, start=1):
            pages += 1
            runs: list[tuple[str, tuple[float, float, float, float], float]] = []

            def visit(
                text: str,
                _cm: Any,
                tm: Any,
                _font: Any,
                size: Any,
                _runs: list[tuple[str, tuple[float, float, float, float], float]] = runs,
            ) -> None:
                cleaned = text.strip()
                if not cleaned:
                    return
                points = float(size or 12.0)
                x, y = float(tm[4]), float(tm[5])
                width = len(cleaned) * points * AVERAGE_GLYPH_WIDTH
                _runs.append((cleaned, (x, y, x + width, y + points), points))

            page.extract_text(visitor_text=visit)
            drafts.extend(self._embedded(page, number=number, asset_path=asset_path))
            for text, bbox, points in runs:
                drafts.append(
                    self.draft(
                        SPAN,
                        {
                            "page": number,
                            "bbox": list(bbox),
                            "bbox_is_estimated": True,
                            "style_hints": {"font_size_pt": points},
                            "asset_path": asset_path,
                        },
                        content_text=text,
                    )
                )
        if not any(draft.kind == SPAN for draft in drafts):
            # The scan-of-a-brochure case. `pypdf` reports success and empty
            # strings, so "no exception" is not evidence that anything was read.
            raise ConnectorDegraded(
                f"the PDF has no text layer across {pages} page(s) — it is probably a scan, "
                "and OCR is not available in this image"
            )
        return drafts

    def _embedded(self, page: Any, *, number: int, asset_path: str) -> list[EvidenceDraft]:
        """Images inside the book, kept as candidate logos.

        A logo lives inside the brand book far more often than it arrives
        beside it, and 3.1.3 cannot propose a clear-space rule for an asset it
        was never shown. Failures here are swallowed on purpose: a page with
        one exotic image filter must not cost the text of the whole document.
        """
        drafts: list[EvidenceDraft] = []
        try:
            embedded = list(page.images)
        except Exception as exc:  # noqa: BLE001 — one odd filter, not a lost document
            log.warning("brand_book.images_unreadable", page=number, error=str(exc))
            return []
        for image in embedded:
            try:
                rendered = image.image
                if rendered is None:
                    continue
                rendered = rendered.convert("RGB")
            except Exception as exc:  # noqa: BLE001 — same reasoning
                log.warning("brand_book.image_undecodable", page=number, error=str(exc))
                continue
            drafts.append(
                self.draft(
                    ASSET,
                    {
                        "page": number,
                        "asset_path": asset_path,
                        "name": getattr(image, "name", None),
                        "width_px": rendered.width,
                        "height_px": rendered.height,
                        "text_layer": "unavailable",
                    },
                    content_text=f"brand asset {rendered.width}x{rendered.height} on page {number}",
                )
            )
            for hex_value, share in _palette(rendered):
                drafts.append(
                    self.draft(
                        COLOUR,
                        {
                            "page": number,
                            "hex": hex_value,
                            "share": share,
                            "asset_path": asset_path,
                        },
                        content_text=hex_value,
                    )
                )
        return drafts

    def _docx(self, content: bytes, *, asset_path: str) -> list[EvidenceDraft]:
        """Paragraph index, not a box.

        A `.docx` stores a flow, not a layout: there is no page and no box
        until Word paginates it, and both depend on the reader's fonts. An
        index points at something a person can find with Ctrl-G; a fabricated
        rectangle points at nothing.
        """
        from docx import Document

        document = Document(io.BytesIO(content))
        drafts: list[EvidenceDraft] = []
        for position, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            drafts.append(
                self.draft(
                    SPAN,
                    {
                        "page": None,
                        "paragraph": position,
                        "bbox": None,
                        "bbox_is_estimated": False,
                        "style_hints": {"style": paragraph.style.name if paragraph.style else None},
                        "asset_path": asset_path,
                    },
                    content_text=text,
                )
            )
        if not drafts:
            raise ConnectorDegraded("the document has no readable paragraphs")
        return drafts

    def _pptx(self, content: bytes, *, asset_path: str) -> list[EvidenceDraft]:
        """The only format that hands over a real box, in EMU, per shape."""
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        drafts: list[EvidenceDraft] = []
        for number, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                left, top = int(shape.left or 0), int(shape.top or 0)
                width, height = int(shape.width or 0), int(shape.height or 0)
                drafts.append(
                    self.draft(
                        SPAN,
                        {
                            "page": number,
                            "bbox": [left, top, left + width, top + height],
                            "bbox_is_estimated": False,
                            "bbox_unit": "emu",
                            "style_hints": {"shape": shape.shape_type and str(shape.shape_type)},
                            "asset_path": asset_path,
                        },
                        content_text=text,
                    )
                )
        if not drafts:
            raise ConnectorDegraded("the deck has no readable text")
        return drafts

    def _image(self, content: bytes, *, asset_path: str) -> list[EvidenceDraft]:
        """A loose image is a candidate logo and a palette. It is never a span.

        No OCR: `tesseract` is not in any image in this repo yet (PRD §21 asks
        for the layer and S3-P0 did not add it). Law 31 says a blocking check
        with no detector is `indeterminate`, never `pass` — so an image
        contributes an asset and its colours, and claims no text at all,
        rather than quietly contributing an empty text layer that would read
        downstream as "the brand book said nothing about this".
        """
        from PIL import Image

        image = Image.open(io.BytesIO(content)).convert("RGB")
        drafts = [
            self.draft(
                ASSET,
                {
                    "page": None,
                    "asset_path": asset_path,
                    "width_px": image.width,
                    "height_px": image.height,
                    "text_layer": "unavailable",
                },
                content_text=f"brand asset {image.width}x{image.height}",
            )
        ]
        for hex_value, share in _palette(image):
            drafts.append(
                self.draft(
                    COLOUR,
                    {"page": None, "hex": hex_value, "share": share, "asset_path": asset_path},
                    content_text=hex_value,
                )
            )
        return drafts

    async def fetch(self, params: dict[str, Any]) -> list[EvidenceDraft]:
        """Every uploaded brand-book file, as spans, assets and colours.

        **Nothing calls this with files yet, and that is a known open gap.**
        §10.2 says brand books arrive "via the existing upload route, stored
        through `storage/backend.py` on the worker Volume" — and that storage
        step does not exist: `POST /projects/{id}/documents` reads the file,
        extracts its text and stores passages, keeping no raw bytes for a
        connector to come back to. So `params["files"]` is never populated, a
        `brand_book` pull returns an empty list, and every guideline run behaves
        as though no brand book had been uploaded: `extraction_confidence: low`,
        rules inferred from live creative, and the G5 card saying so.

        That is the correct *degraded* behaviour, which is why this is a gap
        rather than a break. Closing it needs the upload route to persist the
        original bytes and pass them here — §21 assigns the brand-book upload
        UI to S3-P7, so it is listed as an open question rather than built here
        (§22 scope discipline). The parser and its tests are ready for it.

        Partial success is the normal case — a brand pack is a PDF, a deck and
        a folder of logos, and one of them being a scan is not a reason to
        throw away the other two. So a file that cannot be read adds its name
        to the reason and the rest still come back on the exception, which is
        exactly what `ConnectorDegraded` carries drafts for.
        """
        files = params.get("files") or []
        drafts: list[EvidenceDraft] = []
        failed: list[str] = []
        for entry in files:
            filename = str(entry.get("filename") or "")
            try:
                drafts.extend(
                    self.extract(
                        filename,
                        entry.get("content") or b"",
                        asset_path=str(entry.get("asset_path") or ""),
                    )
                )
            except ConnectorDegraded as exc:
                log.warning("brand_book.file_degraded", filename=filename, reason=exc.reason)
                failed.append(f"{filename}: {exc.reason}")
            except Exception as exc:  # noqa: BLE001 — one bad file, not a failed run
                log.warning("brand_book.file_unreadable", filename=filename, error=str(exc))
                failed.append(f"{filename}: {exc}")
        if failed:
            raise ConnectorDegraded("; ".join(failed), drafts=drafts)
        return drafts


def _palette(image: Any, *, top: int = 5) -> list[tuple[str, float]]:
    """The most common colours, as lowercase hex with the share of pixels each holds.

    Sorted by count and then by hex so that two runs over the same image agree
    on the order of two colours with identical counts. Determinism is not
    decoration here: these tokens reach a ruleset that is hashed.
    """
    counts = image.getcolors(maxcolors=image.width * image.height) or []
    total = float(sum(count for count, _ in counts)) or 1.0
    ranked = sorted(counts, key=lambda item: (-item[0], _hex(item[1])))
    return [(_hex(rgb), round(count / total, 4)) for count, rgb in ranked[:top]]


def _hex(rgb: tuple[int, int, int]) -> str:
    return "#{:02x}{:02x}{:02x}".format(*rgb[:3])
