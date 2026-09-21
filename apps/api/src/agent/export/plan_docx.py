"""Plan DOCX, with a TOC Word fills on open (Stage 02 PRD §14).

The Stage 02 sibling of `export/docx.py`, and it reuses that module's
primitives outright — `_field`, `_table`, `_heading`, `_clear_body` — rather
than restating them. Two implementations of "write a real Word field" would
eventually disagree about `w:dirty`, and the TOC would quietly stop
populating in one of the two documents.

§14's acceptance for this format: "DOCX opens in Word 2019+ and its TOC
populates on F9." That needs real `Heading 1/2/3` styles from
`templates/reference.docx` and a real `fldChar` field marked dirty, both of
which live in `export/docx.py` already.

**The draft watermark is the one thing this file adds to the mechanism.** A
Word document has no `position: fixed`, so the mark cannot be painted once and
repeated the way the PDF does it. It goes in the **header**, which Word repeats
on every page by definition — which is exactly §14's requirement, arrived at by
the one route Word actually offers.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from agent.export.docx import (
    MUTED,
    REFERENCE_DOCX,
    _clear_body,
    _field,
    _heading,
    _mark_fields_dirty,
    _note,
    _page_number_footer,
    _table,
)
from agent.export.plan_contract import CampaignPlan
from agent.export.plan_view import SECTION_TITLES, build_context
from agent.export.templating import fmt_days, fmt_money, fmt_number, fmt_pct, fmt_text

if TYPE_CHECKING:  # pragma: no cover — typing only
    from docx.document import Document as DocxDocument

DRAFT_RED = RGBColor(0xB9, 0x1C, 0x1C)
FROZEN_GREEN = RGBColor(0x04, 0x78, 0x57)


def render_plan_docx(
    plan: CampaignPlan, *, project_name: str | None = None, frozen_by_name: str = ""
) -> bytes:
    """The plan as a Word document."""
    context = build_context(plan, project_name=project_name, frozen_by_name=frozen_by_name)
    document: DocxDocument = Document(str(REFERENCE_DOCX))
    _clear_body(document)
    _status_header(document, plan, context)
    _page_number_footer(document, str(plan.plan_run_id))

    _cover(document, plan, context)
    _toc(document)
    _summary(document, plan, context)
    _objectives(document, plan)
    _media_plan(document, plan, context)
    _channel_slate(document, plan)
    _account_structure(document, plan, context)
    _measurement(document, plan)
    _backlog(document, plan, context)
    _decisions(document, plan, context)
    _dependencies(document, plan, context)

    _mark_fields_dirty(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# chrome
# ---------------------------------------------------------------------------


def _status_header(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    """§14's watermark, by the only mechanism Word repeats on every page."""
    header = document.sections[0].header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if plan.is_frozen:
        run = paragraph.add_run(
            f"FROZEN — v{plan.version} · {plan.generated_at.date().isoformat()}"
        )
        run.font.color.rgb = FROZEN_GREEN
    else:
        run = paragraph.add_run(
            f"{context['watermark']} — figures can still change; not authority to spend"
        )
        run.font.color.rgb = DRAFT_RED
    run.bold = True
    run.font.size = Pt(8)


def _cover(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    title = document.add_heading("Campaign Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph()
    run = subtitle.add_run(context["project_name"])
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED

    _table(
        document,
        ["", ""],
        [
            ["Status", f"{context['status_label']} ({context['version_label']})"],
            ["Plan run", str(plan.plan_run_id)],
            ["Generated", plan.generated_at.strftime("%Y-%m-%d %H:%M UTC")],
            ["Planned from research run", str(plan.source.research_run_id)],
            ["Research accepted", plan.source.accepted_at.date().isoformat()],
            ["Launch readiness at acceptance", fmt_text(plan.source.launch_readiness)],
            ["Constants", fmt_text(plan.constants_version)],
            ["Planning cost", fmt_money(plan.cost_usd, "USD")],
            [
                "Size",
                f"{context['counts']['campaigns']} campaigns · "
                f"{context['counts']['ad_groups']} ad groups · "
                f"{context['counts']['keywords']} keywords",
            ],
        ],
    )
    if plan.source.degraded_sources:
        _note(
            document,
            "Degraded sources: "
            + ", ".join(plan.source.degraded_sources)
            + ". Every forecast in this plan is correspondingly thinner.",
        )
    if plan.source.override_reason:
        _note(document, f"Accepted with an override: {plan.source.override_reason}")


def _toc(document: DocxDocument) -> None:
    heading = document.add_heading("Contents", level=1)
    heading.paragraph_format.page_break_before = True
    _field(document.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ')


# ---------------------------------------------------------------------------
# sections — the same nine, in the same order, as the markdown and the PDF
# ---------------------------------------------------------------------------


def _summary(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    _heading(document, SECTION_TITLES[0], 1)
    document.add_paragraph(plan.executive_summary or "—")

    for issue in context["blocking_issues"]:
        _heading(document, f"Blocking — {fmt_text(issue.get('section'))}", 2)
        document.add_paragraph(fmt_text(issue.get("finding")))
        _note(document, f"Fix. {fmt_text(issue.get('fix'))}")
    for issue in context["warning_issues"]:
        _note(
            document,
            f"Warning — {fmt_text(issue.get('section'))}: {fmt_text(issue.get('finding'))}",
        )


def _objectives(document: DocxDocument, plan: CampaignPlan) -> None:
    objectives = plan.objectives
    _heading(document, SECTION_TITLES[1], 1)
    if objectives.north_star_metric:
        target = (
            f" — {fmt_number(objectives.north_star_target.value, 2)} "
            f"{objectives.north_star_target.unit}"
            if objectives.north_star_target
            else ""
        )
        document.add_paragraph(
            f"North star: {objectives.north_star_metric}{target}, "
            f"{fmt_text(objectives.north_star_period)}."
        )

    _heading(document, "What we can afford to pay", 2)
    _table(
        document,
        ["Segment", "ACV", "Margin", "Lead→won", "Max CPL", "Target CPL", "Max CPA", "Payback"],
        [
            [
                row.segment,
                fmt_money(row.acv_usd, "USD", 0),
                fmt_pct(row.gross_margin_pct),
                fmt_pct(row.lead_to_won_pct, 1),
                fmt_money(row.max_cpl_usd, "USD"),
                fmt_money(row.target_cpl_usd, "USD"),
                fmt_money(row.max_cpa_won_usd, "USD"),
                f"{fmt_number(row.payback_months, 1)} mo",
            ]
            for row in objectives.unit_economics
        ],
    )
    if objectives.method_notes:
        _note(document, objectives.method_notes)

    _heading(document, "What each campaign is aiming at", 2)
    _table(
        document,
        ["Campaign", "Objective", "KPI", "Target", "Ceiling", "Basis"],
        [
            [
                row.campaign_ref,
                row.objective,
                row.primary_kpi,
                fmt_number(row.target_value, 2),
                fmt_number(row.ceiling_value, 2),
                fmt_text(row.basis),
            ]
            for row in objectives.campaign_objectives
        ],
    )

    _heading(document, "What counts as a conversion", 2)
    _table(
        document,
        ["Action", "Category", "Counting", "Value", "Lead→won", "Primary"],
        [
            [
                row.name,
                fmt_text(row.category),
                fmt_text(row.counting),
                fmt_money(row.assigned_value_usd, "USD"),
                fmt_pct(row.lead_to_won_rate_pct, 1),
                "Yes" if row.primary else "No",
            ]
            for row in objectives.conversion_actions
        ],
    )


def _media_plan(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    media = plan.media_plan
    _heading(document, SECTION_TITLES[2], 1)
    if media.envelope is not None:
        currency = media.envelope.currency
        reserve = (
            f" · Experiment reserve {fmt_money(media.experiment_reserve.value, 'USD')}"
            if media.experiment_reserve
            else ""
        )
        document.add_paragraph(
            f"Envelope: {fmt_money(media.envelope.monthly_cap.value, currency)} a month. "
            f"Scenario: {fmt_text(media.chosen_scenario)}.{reserve}"
        )
        document.add_paragraph(fmt_text(media.rationale))
    else:
        _note(document, "No envelope was agreed. This plan states no budget.")

    _heading(document, "Where the money goes", 2)
    _table(
        document,
        ["Campaign", "Market", "Stage", "Monthly", "Share", "Forecast CPA", "Est. conv."],
        [
            [
                row.campaign_ref,
                fmt_text(row.market),
                fmt_text(row.funnel_stage),
                fmt_money(row.usd, "USD", 0),
                fmt_pct(row.pct, 1),
                fmt_money(row.forecast_cpa_usd, "USD"),
                fmt_number(row.est_conv, 1),
            ]
            for row in media.allocation
        ],
    )

    if media.scenarios:
        _heading(document, "The scenarios it was chosen from", 2)
        _table(
            document,
            ["Scenario", "Monthly", "Est. clicks", "Est. conv.", "Est. CPA"],
            [
                [
                    row.name + (" ✓" if row.name == media.chosen_scenario else ""),
                    fmt_money(row.monthly_total_usd, "USD", 0),
                    fmt_number(row.est_clicks),
                    fmt_number(row.est_conv, 1),
                    fmt_money(row.est_cpa, "USD"),
                ]
                for row in media.scenarios
            ],
        )

    if context["forecast"]:
        _heading(document, "The demand it is built on", 2)
        _table(
            document,
            ["Cluster", "Market", "Month", "Impr.", "Clicks", "CPC", "Conv.", "Cost"],
            [
                [
                    fmt_text(row.cluster),
                    fmt_text(row.market),
                    fmt_text(row.month),
                    fmt_number(row.impressions),
                    fmt_number(row.clicks),
                    fmt_money(row.avg_cpc_usd, "USD"),
                    fmt_number(row.conversions, 1),
                    fmt_money(row.cost_usd, "USD", 0),
                ]
                for row in context["forecast"]
            ],
        )
        if context["forecast_omitted"]:
            _note(
                document,
                f"{context['forecast_omitted']} further forecast rows omitted. "
                "The XLSX export carries the whole grid.",
            )


def _channel_slate(document: DocxDocument, plan: CampaignPlan) -> None:
    _heading(document, SECTION_TITLES[3], 1)
    _table(
        document,
        ["Channel", "Market", "Wave", "Share", "Rationale"],
        [
            [
                row.campaign_type,
                fmt_text(row.market),
                fmt_text(row.launch_wave),
                fmt_pct(row.est_share_of_budget_pct, 1),
                fmt_text(row.rationale),
            ]
            for row in plan.channel_slate.slate
        ],
    )
    isolation = plan.channel_slate.brand_isolation
    if isolation is not None:
        _heading(document, "Brand isolation", 2)
        document.add_paragraph(
            f"Brand terms are bid on in {fmt_text(isolation.brand_campaign_ref)} only, at "
            f"{fmt_pct(isolation.budget_pct, 1)} of budget, and are negative everywhere else."
        )
        _note(document, f"Reporting. {fmt_text(isolation.reporting_rule)}")
        _note(document, f"Competitor bidding. {fmt_text(isolation.competitor_bidding_policy)}")


def _account_structure(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    counts = context["counts"]
    _heading(document, SECTION_TITLES[4], 1)
    document.add_paragraph(
        f"{counts['campaigns']} campaigns · {counts['ad_groups']} ad groups · "
        f"{counts['keywords']} keywords."
        + (
            f" Structure verdict: {plan.account_structure.structure_verdict}."
            if plan.account_structure.structure_verdict
            else ""
        )
    )
    for index, campaign in enumerate(context["campaigns"], start=1):
        _heading(document, f"{index}. {campaign.name}", 2)
        _note(
            document,
            f"{fmt_text(campaign.type)} · {fmt_text(campaign.market)} · "
            f"{fmt_money(campaign.daily_budget_usd, 'USD')}/day · "
            f"{fmt_text(campaign.bid_strategy)}",
        )
        _table(
            document,
            ["Ad group", "Theme", "Landing page", "Keywords"],
            [
                [
                    group["name"],
                    fmt_text(group["theme"]),
                    fmt_text(group["landing_url"]),
                    str(group["keyword_count"]),
                ]
                for group in context["ad_group_rows"](campaign)
            ],
        )
    if context["campaigns_omitted"]:
        _note(
            document,
            f"{context['campaigns_omitted']} further campaigns omitted. "
            "The Editor CSV export carries the complete account.",
        )
    if plan.account_structure.account_negatives:
        _heading(document, "Account-level negatives", 2)
        document.add_paragraph(", ".join(plan.account_structure.account_negatives))


def _measurement(document: DocxDocument, plan: CampaignPlan) -> None:
    measurement = plan.measurement_plan
    _heading(document, SECTION_TITLES[5], 1)
    document.add_paragraph(
        f"Source of truth: {fmt_text(measurement.primary_source)}. "
        f"{fmt_text(measurement.rationale)}"
    )
    _table(
        document,
        ["Metric", "How it is computed", "Read from", "Owner", "Refresh"],
        [
            [
                fmt_text(row.get("metric")),
                fmt_text(row.get("formula")),
                fmt_text(row.get("source_field")),
                fmt_text(row.get("owner")),
                fmt_text(row.get("refresh")),
            ]
            for row in measurement.metric_definitions
        ],
    )
    if measurement.reconciliation:
        _heading(document, "Reconciliation", 2)
        _table(
            document,
            ["Metric", "Systems", "Tolerance", "Cadence", "Owner"],
            [
                [
                    fmt_text(row.get("metric")),
                    fmt_text(row.get("systems")),
                    fmt_pct(row.get("tolerance_pct"), 1),
                    fmt_text(row.get("cadence")),
                    fmt_text(row.get("owner")),
                ]
                for row in measurement.reconciliation
            ],
        )
    if measurement.upload:
        _heading(document, "Getting closed deals back into the account", 2)
        document.add_paragraph(
            f"Method: {fmt_text(measurement.upload.get('method'))} · "
            f"Cadence: {fmt_text(measurement.upload.get('cadence'))} · "
            f"Lag: {fmt_days(measurement.upload.get('lag_days'))} · "
            f"Backfill: {fmt_days(measurement.upload.get('backfill_days'))}"
        )
    _heading(document, "Consent", 2)
    document.add_paragraph(
        f"Permitted: {', '.join(measurement.consent_markets_allowed) or '—'}. "
        f"Refused at gate 1.5.3: {', '.join(measurement.consent_markets_blocked) or '—'}."
    )


def _backlog(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    _heading(document, SECTION_TITLES[6], 1)
    if not plan.experiment_backlog:
        document.add_paragraph("No test could be sized against this plan's forecasts.")
        return
    document.add_paragraph(
        f"{len(context['funded_tests'])} of {len(plan.experiment_backlog)} planned tests "
        "are funded from the experiment reserve."
    )
    _table(
        document,
        [
            "#",
            "Test",
            "Campaign",
            "Metric",
            "Conv./arm",
            "Time to read",
            "ICE",
            "Reserve",
            "Funded",
        ],
        [
            [
                row["rank"],
                row["hypothesis"],
                row["campaign"],
                row["metric"],
                row["conv_per_arm"],
                row["days"],
                row["ice"],
                row["reserve"],
                row["funded"],
            ]
            for row in context["test_rows"](plan.experiment_backlog)
        ],
    )


def _decisions(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    _heading(document, SECTION_TITLES[7], 1)
    _table(
        document,
        ["Gate", "Decision", "Status", "Decided by", "Date", "Note"],
        [
            [
                row.gate_key,
                context["gate_labels"].get(row.gate_key, row.name),
                row.status,
                fmt_text(row.decided_by_name),
                row.decided_at.date().isoformat() if row.decided_at else "—",
                fmt_text(row.note),
            ]
            for row in plan.decisions
        ],
    )


def _dependencies(document: DocxDocument, plan: CampaignPlan, context: dict[str, Any]) -> None:
    _heading(document, SECTION_TITLES[8], 1)
    if context["blocking_dependencies"]:
        _heading(document, "Blocking — these stop the plan", 2)
        _table(
            document,
            ["Task", "Owner", "Raised by"],
            [
                [row.task, row.owner, fmt_text(row.source)]
                for row in context["blocking_dependencies"]
            ],
        )
    if context["other_dependencies"]:
        _heading(document, "Everything else", 2)
        for row in context["other_dependencies"]:
            document.add_paragraph(f"{row.task} — {row.owner}", style="List Bullet")
    if not plan.open_dependencies:
        document.add_paragraph("Nothing is outstanding.")

    if plan.assumptions:
        _heading(document, "Assumptions", 2)
        for claim in plan.assumptions:
            document.add_paragraph(
                f"{claim.statement} ({claim.confidence} confidence)", style="List Number"
            )
    if plan.risks:
        _heading(document, "Risks", 2)
        for claim in plan.risks:
            document.add_paragraph(
                f"{claim.statement} ({claim.confidence} confidence)", style="List Number"
            )
    _note(
        document,
        f"{len(plan.numbers())} figures in this plan resolve to a recorded calculation. "
        f"Constants version {fmt_text(plan.constants_version)}.",
    )
