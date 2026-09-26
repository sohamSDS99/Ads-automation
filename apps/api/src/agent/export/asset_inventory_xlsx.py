"""The asset inventory — every text asset and every rendition as a workbook (Stage 04 PRD §14).

Sheet `text`: one row per text asset — id, kind, surface, campaign, ad group,
text, chars / limit, lint verdict, rule ids, claim ids, `ruleset_version`,
`generated_by_ai`, lineage origin. Sheet `media`: one row per rendition —
ratio, px, bytes / limit, derivation, `sx`/`sy`, sha256, model, cost,
disclosure. `chars / limit` and `bytes / limit` are two numeric columns each,
so a person can sort and sum them.

* **Conditional formats are rules, not colours.** Chars at ≥ 90 % of the
  limit and any `pass_with_warnings` verdict are openpyxl `FormulaRule`s over
  the data range, so a cell edited in Excel re-evaluates — the Stage 03
  lesson: a pre-coloured cell satisfies a screenshot and nothing else.
* **Cost is the asset's, on its first rendition only.** Provenance is per
  asset; repeating it on every rendition row would make a column sum count an
  image once per crop.
* **Row 1 says what the file is**, above the header row: a draft's reads
  `DRAFT — NOT RELEASED` (§14: "in the XLSX header row"), and the printed
  page header repeats it; the footer carries the status and creative run id.
* **Deterministic.** Workbook dates pinned to `released_at` (a draft: its
  row's last change), archive stamps normalised by `archives.normalise_zip`.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from decimal import Decimal
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from agent.export.archives import normalise_zip
from agent.export.creative_sources import CreativeExportSources, chars
from agent.export.templating import fmt_datetime

TEXT_COLUMNS: Final = (
    "id",
    "kind",
    "surface",
    "campaign",
    "ad group",
    "text",
    "chars",
    "limit",
    "lint verdict",
    "rule ids",
    "claim ids",
    "ruleset_version",
    "generated_by_ai",
    "lineage origin",
)
MEDIA_COLUMNS: Final = (
    "asset id",
    "media id",
    "campaign",
    "concept",
    "modality",
    "surface",
    "ratio",
    "px",
    "bytes",
    "limit",
    "derivation",
    "sx",
    "sy",
    "sha256",
    "model",
    "cost",
    "disclosure",
    "lint verdict",
    "path",
)
TEXT_WIDTHS: Final = (38, 14, 22, 24, 22, 60, 7, 7, 18, 20, 38, 14, 10, 14)
MEDIA_WIDTHS: Final = (38, 38, 24, 10, 9, 14, 8, 11, 11, 11, 16, 6, 6, 66, 24, 9, 40, 18, 48)

TITLE_ROW: Final = 1
HEADER_ROW: Final = 2
FIRST_DATA_ROW: Final = HEADER_ROW + 1
#: §14: "chars ≥ 90% of limit".
NEAR_LIMIT: Final = 0.9
WARNING: Final = "pass_with_warnings"

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WATERMARK_FONT = Font(color="B91C1C", bold=True, size=12)
TITLE_FONT = Font(bold=True, color="047857")
NEAR_FILL = PatternFill("solid", fgColor="FEF3C7", bgColor="FEF3C7")
WARNING_FILL = PatternFill("solid", fgColor="FFE4E6", bgColor="FFE4E6")


def _cost(value: Decimal | None) -> float | str:
    return float(value) if value is not None else "incomplete"


def _rule(formula: str, fill: PatternFill) -> Rule:
    """`FormulaRule`, typed (openpyxl ships it untyped)."""
    rule: Rule = FormulaRule(formula=[formula], fill=fill)  # type: ignore[no-untyped-call]
    return rule


def _title(sheet: Worksheet, sources: CreativeExportSources) -> None:
    package = sources.package
    header, footer = sheet.oddHeader, sheet.oddFooter
    if header is None or footer is None:  # pragma: no cover - openpyxl creates both
        raise RuntimeError("openpyxl gave the sheet no page header or footer")
    if sources.watermark:
        sheet.cell(row=TITLE_ROW, column=1).value = (
            f"{sources.watermark} · status {package.status} · creative run "
            f"{package.creative_run_id}"
        )
        sheet.cell(row=TITLE_ROW, column=1).font = WATERMARK_FONT
        header.center.text = sources.watermark
    else:
        sheet.cell(row=TITLE_ROW, column=1).value = (
            f"Creative package v{package.version} · {package.status} · released "
            f"{fmt_datetime(sources.released_at)} · package_hash {package.package_hash}"
        )
        sheet.cell(row=TITLE_ROW, column=1).font = TITLE_FONT
    footer.left.text = f"Status {package.status} · creative run {package.creative_run_id}"
    footer.right.text = "Page &P of &N"


def _table(
    sheet: Worksheet,
    columns: Sequence[str],
    widths: Sequence[int],
    rows: Sequence[Sequence[Any]],
) -> None:
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    for offset, row in enumerate(rows):
        for index, value in enumerate(row, start=1):
            sheet.cell(row=FIRST_DATA_ROW + offset, column=index, value=value)
    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    last = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last}{HEADER_ROW + max(len(rows), 1)}"


def _data_range(sheet: Worksheet, columns: Sequence[str], rows: int) -> str:
    return f"A{FIRST_DATA_ROW}:{get_column_letter(len(columns))}{FIRST_DATA_ROW + max(rows, 1) - 1}"


def _column(columns: Sequence[str], name: str) -> str:
    return get_column_letter(columns.index(name) + 1)


def _text_rows(sources: CreativeExportSources) -> list[list[Any]]:
    rows = []
    for campaign in sources.package.campaigns:
        name = sources.campaigns.get(campaign.campaign_ref)
        for asset in campaign.text_assets:
            rows.append(
                [
                    str(asset.asset_id),
                    asset.kind,
                    asset.surface,
                    name.name if name is not None else campaign.campaign_ref,
                    asset.ad_group_ref,
                    asset.text,
                    chars(asset.text) if asset.text is not None else None,
                    sources.limit(campaign.campaign_type, asset.surface),
                    asset.lint.verdict,
                    ", ".join(asset.lint.rule_ids) or None,
                    ", ".join(str(claim) for claim in asset.claim_ids) or None,
                    asset.ruleset_version,
                    asset.generated_by_ai,
                    asset.lineage.get("origin"),
                ]
            )
    return rows


def _media_rows(sources: CreativeExportSources) -> list[list[Any]]:
    rows = []
    for campaign in sources.package.campaigns:
        name = sources.campaigns.get(campaign.campaign_ref)
        for asset in (*campaign.media, *campaign.logos):
            for index, rendition in enumerate(asset.renditions):
                rows.append(
                    [
                        str(asset.asset_id),
                        str(rendition.media_id),
                        name.name if name is not None else campaign.campaign_ref,
                        asset.concept_id,
                        asset.modality,
                        rendition.surface,
                        rendition.aspect_ratio,
                        f"{rendition.width}x{rendition.height}",
                        rendition.bytes,
                        sources.max_bytes(
                            campaign.campaign_type, asset.modality, rendition.aspect_ratio
                        ),
                        rendition.derivation,
                        rendition.scale[0],
                        rendition.scale[1],
                        rendition.sha256,
                        asset.provenance.model_id,
                        _cost(asset.provenance.cost_usd) if index == 0 else None,
                        ", ".join(
                            f"{key}: {value}"
                            for key, value in sorted((rendition.disclosure or {}).items())
                        )
                        or None,
                        rendition.lint.verdict,
                        rendition.path,
                    ]
                )
    return rows


def render_asset_inventory_xlsx(sources: CreativeExportSources) -> bytes:
    book = Workbook()
    # `Workbook()` stamps both dates from the wall clock; the document's own
    # time replaces them (see `budget_xlsx`).
    book.properties.created = sources.stamped
    book.properties.modified = sources.stamped
    book.properties.title = f"Asset inventory — {sources.project_name or 'creative package'}"
    book.remove(book.active)  # type: ignore[arg-type]

    text = book.create_sheet("text")
    text_rows = _text_rows(sources)
    _title(text, sources)
    _table(text, TEXT_COLUMNS, TEXT_WIDTHS, text_rows)
    first = FIRST_DATA_ROW
    count, limit, verdict = (_column(TEXT_COLUMNS, n) for n in ("chars", "limit", "lint verdict"))
    cells = _data_range(text, TEXT_COLUMNS, len(text_rows))
    near = f"AND(ISNUMBER(${limit}{first}),${count}{first}>={NEAR_LIMIT}*${limit}{first})"
    text.conditional_formatting.add(cells, _rule(near, NEAR_FILL))
    text.conditional_formatting.add(cells, _rule(f'${verdict}{first}="{WARNING}"', WARNING_FILL))

    media = book.create_sheet("media")
    media_rows = _media_rows(sources)
    _title(media, sources)
    _table(media, MEDIA_COLUMNS, MEDIA_WIDTHS, media_rows)
    media_verdict = _column(MEDIA_COLUMNS, "lint verdict")
    media.conditional_formatting.add(
        _data_range(media, MEDIA_COLUMNS, len(media_rows)),
        _rule(f'${media_verdict}{first}="{WARNING}"', WARNING_FILL),
    )

    buffer = io.BytesIO()
    book.save(buffer)
    return normalise_zip(buffer.getvalue(), sources.stamped)
