"""The asset inventory XLSX (Stage 04 PRD §14).

Read back with openpyxl, never trusted: the columns are looked up by header,
the conditional formats are read off the sheet (a real rule, not a
pre-coloured cell — the Stage 03 lesson), and determinism is checked under a
forced clock shift, because two saves inside one second match whatever the
archive carries.
"""

from __future__ import annotations

import io
import re
import time
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from agent.export.asset_inventory_xlsx import (
    MEDIA_COLUMNS,
    TEXT_COLUMNS,
    render_asset_inventory_xlsx,
)
from agent.export.creative_sources import DRAFT_WATERMARK
from tests.creative import export_support as support
from tests.creative import package_support as golden

HEADER_ROW = 2


def _book(blob: bytes) -> Any:
    return load_workbook(io.BytesIO(blob))


def _rows(sheet: Worksheet) -> list[dict[str, Any]]:
    header = [cell.value for cell in sheet[HEADER_ROW]]
    return [
        dict(zip(header, (cell.value for cell in row), strict=True))
        for row in sheet.iter_rows(min_row=HEADER_ROW + 1)
    ]


def _letter(sheet: Worksheet, name: str) -> str:
    for cell in sheet[HEADER_ROW]:
        if cell.value == name:
            return str(cell.column_letter)
    raise AssertionError(f"no column {name}")


def _rules(sheet: Worksheet) -> list[tuple[str, str]]:
    return [
        (str(ranges.sqref), formula)
        for ranges in sheet.conditional_formatting
        for rule in ranges.rules
        for formula in rule.formula
    ]


def test_there_are_two_sheets_text_and_media() -> None:
    assert _book(render_asset_inventory_xlsx(support.sources())).sheetnames == ["text", "media"]


def test_the_text_sheet_has_one_row_per_text_asset_with_section_14_columns() -> None:
    sources = support.sources()
    sheet = _book(render_asset_inventory_xlsx(sources))["text"]
    assert [cell.value for cell in sheet[HEADER_ROW]] == list(TEXT_COLUMNS)
    assert list(TEXT_COLUMNS) == [
        "id", "kind", "surface", "campaign", "ad group", "text", "chars", "limit",
        "lint verdict", "rule ids", "claim ids", "ruleset_version", "generated_by_ai",
        "lineage origin",
    ]  # fmt: skip
    rows = _rows(sheet)
    assets = [a for c in sources.package.campaigns for a in c.text_assets]
    assert len(rows) == len(assets)
    first = next(r for r in rows if r["id"] == str(golden.headline_ids("A")[0]))
    assert first["kind"] == "headline" and first["surface"] == "rsa_headline"
    assert first["campaign"] == support.SEARCH_NAME and first["ad group"] == golden.AD_GROUP
    assert (first["chars"], first["limit"]) == (13, 30)
    assert first["lint verdict"] == "pass" and first["ruleset_version"] == golden.PIN
    assert first["generated_by_ai"] is True
    description = next(r for r in rows if r["id"] == str(golden.description_ids("A")[0]))
    assert description["claim ids"] == str(support.LICENSED)
    pmax = next(r for r in rows if r["id"] == str(support.PMAX_LONG))
    assert pmax["lineage origin"] == "generated" and pmax["limit"] == 90


def test_the_media_sheet_has_one_row_per_rendition_with_section_14_columns() -> None:
    sources = support.sources()
    sheet = _book(render_asset_inventory_xlsx(sources))["media"]
    header = [cell.value for cell in sheet[HEADER_ROW]]
    assert header == list(MEDIA_COLUMNS)
    for name in ("ratio", "px", "bytes", "limit", "derivation", "sx", "sy", "sha256", "model",
                 "cost", "disclosure"):  # fmt: skip
        assert name in header, name
    rows = _rows(sheet)
    assert len(rows) == 5
    wide = next(r for r in rows if r["media id"] == str(support.PMAX_IMAGE_WIDE))
    assert wide["ratio"] == "1.91:1" and wide["px"] == "1200x628"
    assert wide["limit"] == 5_242_880 and isinstance(wide["bytes"], int)
    assert (wide["sx"], wide["sy"]) == (0.5, 0.5)
    assert wide["model"] == "openai/gpt-image-1" and wide["cost"] == 0.4
    assert "trainedAlgorithmicMedia" in wide["disclosure"]
    assert len(wide["sha256"]) == 64


def test_chars_at_ninety_percent_of_the_limit_and_warnings_are_conditionally_formatted() -> None:
    book = _book(render_asset_inventory_xlsx(support.sources()))
    text = book["text"]
    chars, limit, verdict = (_letter(text, name) for name in ("chars", "limit", "lint verdict"))
    rules = _rules(text)
    last = text.max_row
    near = [(r, f) for r, f in rules if f"{chars}{HEADER_ROW + 1}>=0.9*" in f.replace("$", "")]
    assert near, rules
    (sqref, formula), *_ = near
    assert f"{limit}{HEADER_ROW + 1}" in formula.replace("$", "")
    assert sqref.endswith(str(last)) and sqref.startswith(f"A{HEADER_ROW + 1}")
    assert any(
        f'{verdict}{HEADER_ROW + 1}="pass_with_warnings"' in f.replace("$", "") for _, f in rules
    )
    media = book["media"]
    media_verdict = _letter(media, "lint verdict")
    assert any(
        f'{media_verdict}{HEADER_ROW + 1}="pass_with_warnings"' in f.replace("$", "")
        for _, f in _rules(media)
    )


def test_a_draft_is_watermarked_in_the_header_row_of_both_sheets() -> None:
    sources = support.draft_sources()
    book = _book(render_asset_inventory_xlsx(sources))
    for sheet in book.worksheets:
        assert str(sheet["A1"].value).startswith(DRAFT_WATERMARK), sheet.title
        assert DRAFT_WATERMARK in (sheet.oddHeader.center.text or "")
        footer = sheet.oddFooter.left.text or ""
        assert "ready_to_release" in footer and str(sources.package.creative_run_id) in footer


def test_a_released_inventory_is_not_watermarked() -> None:
    sources = support.sources()
    book = _book(render_asset_inventory_xlsx(sources))
    for sheet in book.worksheets:
        assert DRAFT_WATERMARK not in str(sheet["A1"].value)
        assert DRAFT_WATERMARK not in (sheet.oddHeader.center.text or "")
        assert "released" in (sheet.oddFooter.left.text or "")


def test_two_exports_are_byte_identical_across_a_clock_shift(monkeypatch: Any) -> None:
    first = render_asset_inventory_xlsx(support.sources())
    real = time.time
    monkeypatch.setattr(time, "time", lambda: real() + 86_400 * 5 + 11)
    assert render_asset_inventory_xlsx(support.sources()) == first


def test_the_document_dates_are_the_release() -> None:
    blob = render_asset_inventory_xlsx(support.sources())
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    stamps = re.findall(r"<dcterms:(created|modified)[^>]*>([^<]*)<", core)
    assert dict(stamps) == {"created": "2026-09-26T15:30:00Z", "modified": "2026-09-26T15:30:00Z"}
