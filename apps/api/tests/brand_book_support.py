"""Real brand-book files, built in memory, for the `brand_book` connector suite.

Hand-built rather than rendered, for the same reason `tests/test_documents.py`
builds its PDFs by hand: WeasyPrint needs pango and cairo, which only the worker
image carries, and this suite has to run on a laptop.

Every builder takes positions as arguments, because the thing under test is
*provenance* — a span that cannot say which page and which box it came from is
not evidence a brand owner can check at gate G5.
"""

from __future__ import annotations

import io
import zlib

#: A string that exists only inside a brand book's bytes. The canary test
#: (PRD §21) asserts it reaches zero prompts and zero payload fields, so it has
#: to be distinctive enough that a substring search cannot match it by accident.
CANARY = "ZZQX-CANARY-BRANDBOOK-8842"


def pdf_bytes(pages: list[list[tuple[str, int, int]]]) -> bytes:
    """A genuine PDF. Each page is a list of `(text, x, y)` text runs.

    Positions are PDF user-space points with the origin at the bottom left,
    which is what the text matrix reports back and therefore what the
    connector has to translate.
    """
    objects: list[bytes] = []
    page_ids = [4 + index * 2 for index, _ in enumerate(pages)]

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for index, runs in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> "
                f"/MediaBox [0 0 612 792] /Contents {content_id} 0 R >>"
            ).encode()
        )
        parts = []
        for text, x, y in runs:
            escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            parts.append(f"BT /F1 12 Tf {x} {y} Td ({escaped}) Tj ET")
        stream = "\n".join(parts).encode()
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )

    return _assemble(objects)


def image_only_pdf() -> bytes:
    """One page carrying an image XObject and no text operators at all.

    The scanned-brochure case. `pypdf` reports success and zero characters, so
    a connector that trusted "no exception" would emit nothing and say nothing.
    """
    pixels = zlib.compress(b"\xff\xff\xff" * 16)
    return _assemble(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /XObject << /Im0 5 0 R >> >> >>",
            b"<< /Length 24 >>\nstream\nq 612 0 0 792 0 0 cm Q\nendstream",
            b"<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceRGB "
            b"/BitsPerComponent 8 /Filter /FlateDecode /Length "
            + str(len(pixels)).encode()
            + b" >>\nstream\n"
            + pixels
            + b"\nendstream",
        ]
    )


def _assemble(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    )
    return bytes(out)


def docx_bytes(paragraphs: list[str]) -> bytes:
    """A real .docx. No geometry: Word has no page box until it is laid out."""
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(slides: list[list[tuple[str, int, int, int, int]]]) -> bytes:
    """A real .pptx. Each slide is a list of `(text, left, top, width, height)` in EMU.

    The one format whose own object model carries a true bounding box per
    shape, which is why it is worth supporting rather than asking people to
    export to PDF.
    """
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for shapes in slides:
        slide = presentation.slides.add_slide(blank)
        for text, left, top, width, height in shapes:
            box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
            box.text_frame.text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def png_bytes(colour: tuple[int, int, int], size: tuple[int, int] = (24, 24)) -> bytes:
    """A flat PNG in one colour — the simplest thing with a palette to sample."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()
