"""S4-P17 — creative-package exports, end to end (Stage 04 PRD §14).

The golden run (S4-P16's: S4-P8's extras, fresh offers, a crawled site, a
landing patch) through the real route, the real worker job and the real
download. The unit suites own the formats' contents; this one owns what only
the database can show: the 409 before anything is queued, names read from
the run's pinned plan, spec limits from the final pin, files read back from
where release wrote them, and byte-identical downloads of one release.
"""

from __future__ import annotations

import csv
import io
import uuid
import zipfile
from typing import Any

import pytest
import sqlalchemy as sa
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from agent.creative.package import package_hash
from agent.db.models import Export, Run
from agent.export.creative_sources import DRAFT_WATERMARK
from agent.export.editor_zip import load_editor_columns
from agent.guardrails.matchers.assets import measure
from agent.schemas.creative_input import CreativeInput
from tests.integration.conftest import ApiClient
from tests.integration.test_s4p8_extras import _Web, web  # noqa: F401 — fixture
from tests.integration.test_s4p16_package import (  # noqa: F401 — fixtures
    _package_of,
    _release,
    api_dependency_overrides,
    golden,
    storage,
    worker_files,
    worker_writes,
)

pytestmark = pytest.mark.usefixtures("web", "worker_writes")


async def _export(api: ApiClient, package_id: uuid.UUID, fmt: str) -> bytes:
    from agent.export.jobs import generate_export

    accepted = await api.post(f"/creative-packages/{package_id}/export", params={"format": fmt})
    assert accepted.status_code == 202, accepted.text
    job_id = accepted.json()["job_id"]
    done = await generate_export({}, job_id)
    assert done["status"] == "ready", done
    download = await api.get(f"/exports/{job_id}/download")
    assert download.status_code == 200, download.text
    return download.content


async def _plan_names(db: AsyncSession, run_id: uuid.UUID) -> dict[str, set[str]]:
    run = await db.get(Run, run_id)
    assert run is not None
    inp = CreativeInput.model_validate(run.creative_input)
    return {
        campaign.name: {group.name for group in campaign.ad_groups}
        for campaign in inp.account_structure.campaigns
    }


async def _golden(admin: ApiClient, db: AsyncSession, ws: uuid.UUID, project: uuid.UUID,
                  actor: uuid.UUID) -> Any:  # fmt: skip
    run_id = await golden(admin, db, ws, project, actor)
    return run_id, await _package_of(db, run_id)


async def test_an_editor_zip_of_an_unreleased_package_is_409_and_queues_nothing(
    admin: ApiClient,
    db: AsyncSession,
    workspace_id: uuid.UUID,
    project_id: uuid.UUID,
    admin_user: Any,
) -> None:
    _, row = await _golden(admin, db, workspace_id, project_id, admin_user.id)
    assert row.status.value == "ready_to_release"
    before = await db.scalar(sa.select(sa.func.count()).select_from(Export))
    refused = await admin.post(
        f"/creative-packages/{row.id}/export", params={"format": "editor_zip"}
    )
    assert refused.status_code == 409, refused.text
    body = refused.json()
    assert body["code"] == "package_not_released"
    assert "ready to release" in body["detail"]
    assert await db.scalar(sa.select(sa.func.count()).select_from(Export)) == before


async def test_a_released_package_exports_an_editor_zip_keyed_by_the_pinned_plan(
    admin: ApiClient,
    db: AsyncSession,
    workspace_id: uuid.UUID,
    project_id: uuid.UUID,
    admin_user: Any,
) -> None:
    run_id, row = await _golden(admin, db, workspace_id, project_id, admin_user.id)
    assert (await _release(admin, row.id, 1)).status_code == 200
    first = await _export(admin, row.id, "editor_zip")
    second = await _export(admin, row.id, "editor_zip")
    assert first == second  # acceptance 4

    columns = load_editor_columns()
    known = {spec.filename: list(spec.columns) for spec in columns.files.values()}
    names = await _plan_names(db, run_id)
    with zipfile.ZipFile(io.BytesIO(first)) as archive:
        entries = archive.namelist()
        assert entries == sorted(entries)
        assert "README.md" in entries and "responsive_search_ads.csv" in entries
        referenced: set[str] = set()
        for name in (n for n in entries if n.endswith(".csv")):
            text = archive.read(name).decode(columns.encoding)
            header, *rows = list(csv.reader(io.StringIO(text)))
            assert header == known[name], name  # acceptance 2: every column is in the yaml
            for values in rows:
                row_map = dict(zip(header, values, strict=True))
                assert row_map["Campaign"] in names, row_map["Campaign"]
                group = row_map.get("Ad group") or row_map.get("Asset group")
                if group:
                    assert group in names[row_map["Campaign"]], group
                referenced |= {v for v in values if v.startswith("media/")}
        assert referenced <= set(entries)  # acceptance 2: every media path exists
        rsa = archive.read("responsive_search_ads.csv").decode(columns.encoding)
    header, *rows = list(csv.reader(io.StringIO(rsa)))
    assert rows, "the golden run ships responsive search ads"
    for values in rows:  # acceptance 2: within the final pin's limits (30 / 90 / 15)
        for column, value in zip(header, values, strict=True):
            limit = 30 if column.startswith("Headline ") and "position" not in column else (
                90 if column.startswith("Description ") and "position" not in column else (
                    15 if column.startswith("Path ") else None))  # fmt: skip
            if limit is not None and value:
                assert measure(value, "chars") <= limit, (column, value)


async def test_the_creative_book_is_watermarked_until_release_and_identical_after(
    admin: ApiClient,
    db: AsyncSession,
    workspace_id: uuid.UUID,
    project_id: uuid.UUID,
    admin_user: Any,
) -> None:
    _, row = await _golden(admin, db, workspace_id, project_id, admin_user.id)
    draft = PdfReader(io.BytesIO(await _export(admin, row.id, "pdf")))
    assert all(DRAFT_WATERMARK in (page.extract_text() or "") for page in draft.pages)
    assert all(str(row.creative_run_id) in (page.extract_text() or "") for page in draft.pages)

    assert (await _release(admin, row.id, 1)).status_code == 200
    await db.refresh(row)
    first = await _export(admin, row.id, "pdf")
    assert first == await _export(admin, row.id, "pdf")
    released = PdfReader(io.BytesIO(first))
    assert not any(DRAFT_WATERMARK in (page.extract_text() or "") for page in released.pages)
    assert row.released_at is not None
    stamp = row.released_at.strftime("D:%Y%m%d%H%M%SZ")
    assert released.metadata is not None and released.metadata["/CreationDate"] == stamp
    text = "\n".join(page.extract_text() or "" for page in released.pages)
    for section in ("1. Brief", "2. Ads by ad group", "4. Media contact sheet", "8. Sign-off"):
        assert section in text


async def test_xlsx_and_md_export_and_the_json_still_verifies(
    admin: ApiClient,
    db: AsyncSession,
    workspace_id: uuid.UUID,
    project_id: uuid.UUID,
    admin_user: Any,
) -> None:
    import json

    from openpyxl import load_workbook

    _, row = await _golden(admin, db, workspace_id, project_id, admin_user.id)
    draft = load_workbook(io.BytesIO(await _export(admin, row.id, "xlsx")))
    assert str(draft["text"]["A1"].value).startswith(DRAFT_WATERMARK)
    assert (await _release(admin, row.id, 1)).status_code == 200
    await db.refresh(row)
    xlsx = await _export(admin, row.id, "xlsx")
    assert xlsx == await _export(admin, row.id, "xlsx")
    book = load_workbook(io.BytesIO(xlsx))
    assert book.sheetnames == ["text", "media"] and book["text"].max_row > 3
    md = (await _export(admin, row.id, "md")).decode("utf-8")
    assert md.startswith("# Creative book") and "## 8. Sign-off" in md
    exported = json.loads(await _export(admin, row.id, "json"))
    assert package_hash(exported) == row.package_hash  # acceptance 1


async def test_a_format_no_package_has_is_still_a_422(
    admin: ApiClient,
    db: AsyncSession,
    workspace_id: uuid.UUID,
    project_id: uuid.UUID,
    admin_user: Any,
) -> None:
    _, row = await _golden(admin, db, workspace_id, project_id, admin_user.id)
    refused = await admin.post(f"/creative-packages/{row.id}/export", params={"format": "docx"})
    assert refused.status_code == 422, refused.text
